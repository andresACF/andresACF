#!/usr/bin/env python3
"""Generate Andrés Cornejo CV (DOCX + PDF) with updated experience."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm, mm
from reportlab.platypus import (
    Frame,
    PageTemplate,
    BaseDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
    KeepTogether,
    ListFlowable,
    ListItem,
)

OUT_DIR = Path(__file__).resolve().parent

# Experience updates:
# - Left Ripio (Enero 2026 – Agosto 2026)
# - Started Soyoda this month as Programador de Sistemas (Septiembre 2026 – Ahora)

ABOUT = (
    "¡Hola! Soy Ingeniero de Ciencias de la Computación en ESPOL, "
    "apasionado por la automatización de procesos y la elaboración "
    "de flujos agénticos. He competido en retos universitarios de "
    "seguridad informática e inteligencia artificial, y poseo experiencia "
    "en desarrollo fullstack, AI Engineering, Data Engineering y "
    "programación de sistemas. Estoy entusiasmado por continuar creciendo "
    "en mi trayectoria profesional y comprometido a dedicarme plenamente en ello."
)

EXPERIENCES = [
    {
        "dates": "Septiembre 2026 - Ahora",
        "company": "SOYODA",
        "role": "Programador de Sistemas",
        "bullets": [
            "Desarrollo y mantenimiento de sistemas de software a medida.",
            "Soporte a la operación y evolución de aplicaciones internas.",
            "Colaboración en análisis, implementación y mejora continua de procesos.",
        ],
    },
    {
        "dates": "Enero 2026 - Agosto 2026",
        "company": "RIPIO",
        "role": "FullStack web developer (Django + React)",
        "bullets": [
            "Fullstack developer (.NET + Django + PostgreSQL).",
            "N8N + desarrollo de agentes de IA.",
            "Web Scraping, RPA.",
            "LangGraph + Redis + OpenViking (Agent memory).",
            "Docker + Caddy, Celery.",
        ],
    },
    {
        "dates": "Junio 2025 - Ahora",
        "company": "Consultor Independiente / Desarrollador Fullstack",
        "role": None,
        "bullets": [
            "MMN (Sistema de Control de Reparaciones).",
            "PRICOTERCORP (Ecommerce) Contifico + Datafast.",
        ],
    },
    {
        "dates": "Mayo 2025 - Noviembre 2025",
        "company": "Banco Guayaquil",
        "role": None,
        "bullets": [
            "Manejo del portal Azure (asignación de roles según matriz RBAC, gestión de grupos de administración).",
            "Active Directory, Trend Vision One y Mosyle.",
            "Power Platform (Power Automate, Power Virtual Agents).",
            "Reducción del 40% en tiempo de asignación de licencias de Teams y Microsoft 365 mediante flujos de Power Automate.",
        ],
    },
    {
        "dates": "Mayo 2024 - Febrero 2025",
        "company": "ESPOL",
        "role": "Desarrollador Backend y Frontend (React + Flutter)",
        "bullets": [
            "Desarrollo Backend y Frontend con React y Flutter.",
            "Mantenimiento de hardware.",
            "API REST.",
        ],
    },
]

EDUCATION = [
    {
        "dates": "2021-2026",
        "title": "Ingeniero en Ciencias de la Computación",
        "place": "ESPOL",
    },
    {
        "dates": "2015-2021",
        "title": "Bachillerato Internacional",
        "place": 'Colegio "San José La Salle - Ecuador"',
    },
]

COURSES = [
    "Graduado de la Academia AWS - Ingeniería de Datos - Insignia de capacitación",
    "Microsoft Word and Excel advanced",
    "Power Point",
]

LANGUAGES = [
    "Español nativo",
    "B2 English — Centro Ecuatoriano Norteamericano",
]

SKILLS = [
    "Atención y soporte al usuario",
    "Power Platforms (Power Automate, Power BI, Power Virtual Agents) + N8N",
    "Lenguajes de programación: Python, Dart, JavaScript, Ruby, PHP, R",
    "Frameworks y herramientas: Django, Flutter, Ruby on Rails, Jenkins, jQuery",
    "Bases de datos: PostgreSQL, Redis, BD tipo RAG para agentes (OpenViking)",
    "Análisis de datos y generación de reportes: Pandas, R, ReportLab (PDF), Web Scraping, análisis de comunidades, modelos de regresión y clustering",
    "DevOps e Integración/Entrega continua (CI/CD): Jenkins, GitHub, Despliegue en AWS",
    "Enfoque en la hiper-automatización de procesos bancarios y corporativos mediante RPA y desarrollo de software a medida",
    "Manejo de Power BI en la inteligencia del negocio",
]

REFERENCES = [
    {
        "name": "Andrés David Loor Villamar",
        "role": "Cyber Identity Engineer",
        "org": "Banco Guayaquil",
        "phone": "+593 97 913 0413",
    },
    {
        "name": "Brian Huber",
        "role": "Manager",
        "org": "Electric Tour Company",
        "phone": "+1 415 474 3130",
    },
    {
        "name": "Ing. Daniel Castro",
        "role": "Asistent DST Area",
        "org": "ESPOL",
        "phone": "+593 98 519 4690",
    },
    {
        "name": "Joffre Morales",
        "role": "Gerente General",
        "org": "Mágico Mundo del Nintendo",
        "phone": "+593 98 718 1478",
    },
]

CONTACT = {
    "phone": "+593 98 359 2954",
    "email": "andrescoraf@gmail.com",
    "email2": "andalcor@espol.edu.ec",
}

NAVY = HexColor("#1B2A4A")
ACCENT = HexColor("#2C4A7C")
LIGHT = HexColor("#EEF2F7")
MUTED = HexColor("#4A5568")
DARK = HexColor("#1A202C")


def build_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("ANDRES ALFREDO\nCORNEJO FIGUEROA")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INGENIERO EN CIENCIAS DE LA COMPUTACIÓN")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x2C, 0x4A, 0x7C)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run(
        f"{CONTACT['phone']}  ·  {CONTACT['email']}  ·  {CONTACT['email2']}"
    ).font.size = Pt(9)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    heading("ACERCA DE MÍ")
    about = doc.add_paragraph(ABOUT)
    about.paragraph_format.space_after = Pt(6)

    heading("EXPERIENCIA")
    for exp in EXPERIENCES:
        p = doc.add_paragraph()
        r = p.add_run(f"{exp['dates']}  |  {exp['company']}")
        r.bold = True
        r.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        if exp["role"]:
            role_p = doc.add_paragraph()
            rr = role_p.add_run(exp["role"])
            rr.italic = True
            rr.font.size = Pt(10)
            role_p.paragraph_format.space_after = Pt(2)
        for bullet in exp["bullets"]:
            bp = doc.add_paragraph(bullet, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)

    heading("EDUCACIÓN")
    for edu in EDUCATION:
        p = doc.add_paragraph()
        r = p.add_run(f"{edu['dates']}  |  {edu['title']}")
        r.bold = True
        doc.add_paragraph(edu["place"])

    heading("CURSOS")
    for course in COURSES:
        doc.add_paragraph(course, style="List Bullet")

    heading("LENGUAJES")
    for lang in LANGUAGES:
        doc.add_paragraph(lang, style="List Bullet")

    heading("HABILIDADES")
    for skill in SKILLS:
        doc.add_paragraph(skill, style="List Bullet")

    heading("REFERENCIAS PERSONALES")
    for ref in REFERENCES:
        p = doc.add_paragraph()
        r = p.add_run(ref["name"])
        r.bold = True
        doc.add_paragraph(f"{ref['role']} — {ref['org']}")
        doc.add_paragraph(ref["phone"])

    doc.save(path)


def build_pdf(path: Path) -> None:
    doc = BaseDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=1.6 * cm,
        rightMargin=1.6 * cm,
        topMargin=1.4 * cm,
        bottomMargin=1.4 * cm,
    )

    styles = {
        "name": ParagraphStyle(
            "name",
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            textColor=NAVY,
            alignment=1,
            spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=13,
            textColor=ACCENT,
            alignment=1,
            spaceAfter=4,
        ),
        "contact": ParagraphStyle(
            "contact",
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=MUTED,
            alignment=1,
            spaceAfter=8,
        ),
        "h1": ParagraphStyle(
            "h1",
            fontName="Helvetica-Bold",
            fontSize=11.5,
            leading=14,
            textColor=NAVY,
            spaceBefore=10,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=DARK,
            spaceAfter=4,
            alignment=4,
        ),
        "job": ParagraphStyle(
            "job",
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=12,
            textColor=DARK,
            spaceBefore=6,
            spaceAfter=1,
        ),
        "role": ParagraphStyle(
            "role",
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=11,
            textColor=ACCENT,
            spaceAfter=2,
        ),
        "bullet": ParagraphStyle(
            "bullet",
            fontName="Helvetica",
            fontSize=8.8,
            leading=11.5,
            textColor=DARK,
            leftIndent=8,
        ),
        "ref_name": ParagraphStyle(
            "ref_name",
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            textColor=DARK,
            spaceBefore=4,
        ),
        "ref_meta": ParagraphStyle(
            "ref_meta",
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=MUTED,
        ),
    }

    frame = Frame(
        doc.leftMargin,
        doc.bottomMargin,
        doc.width,
        doc.height,
        id="normal",
    )
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame])])

    story = []
    story.append(Paragraph("ANDRES ALFREDO<br/>CORNEJO FIGUEROA", styles["name"]))
    story.append(
        Paragraph("INGENIERO EN CIENCIAS DE LA COMPUTACIÓN", styles["subtitle"])
    )
    story.append(
        Paragraph(
            f"{CONTACT['phone']}  ·  {CONTACT['email']}  ·  {CONTACT['email2']}",
            styles["contact"],
        )
    )
    story.append(
        HRFlowable(width="100%", thickness=1.2, color=NAVY, spaceAfter=6)
    )

    story.append(Paragraph("ACERCA DE MÍ", styles["h1"]))
    story.append(Paragraph(ABOUT, styles["body"]))

    story.append(Paragraph("EXPERIENCIA", styles["h1"]))
    for exp in EXPERIENCES:
        block = [
            Paragraph(f"{exp['dates']}  |  {exp['company']}", styles["job"]),
        ]
        if exp["role"]:
            block.append(Paragraph(exp["role"], styles["role"]))
        items = [
            ListItem(Paragraph(b, styles["bullet"]), leftIndent=12, bulletColor=ACCENT)
            for b in exp["bullets"]
        ]
        block.append(
            ListFlowable(
                items,
                bulletType="bullet",
                start="•",
                leftIndent=10,
                bulletFontSize=8,
                spaceBefore=0,
                spaceAfter=2,
            )
        )
        story.append(KeepTogether(block))

    story.append(Paragraph("EDUCACIÓN", styles["h1"]))
    for edu in EDUCATION:
        story.append(
            Paragraph(f"<b>{edu['dates']}</b>  |  {edu['title']}", styles["body"])
        )
        story.append(Paragraph(edu["place"], styles["ref_meta"]))

    story.append(Paragraph("CURSOS", styles["h1"]))
    story.append(
        ListFlowable(
            [
                ListItem(Paragraph(c, styles["bullet"]), leftIndent=12, bulletColor=ACCENT)
                for c in COURSES
            ],
            bulletType="bullet",
            start="•",
            leftIndent=10,
        )
    )

    story.append(Paragraph("LENGUAJES", styles["h1"]))
    story.append(
        ListFlowable(
            [
                ListItem(Paragraph(l, styles["bullet"]), leftIndent=12, bulletColor=ACCENT)
                for l in LANGUAGES
            ],
            bulletType="bullet",
            start="•",
            leftIndent=10,
        )
    )

    story.append(Paragraph("HABILIDADES", styles["h1"]))
    story.append(
        ListFlowable(
            [
                ListItem(Paragraph(s, styles["bullet"]), leftIndent=12, bulletColor=ACCENT)
                for s in SKILLS
            ],
            bulletType="bullet",
            start="•",
            leftIndent=10,
        )
    )

    story.append(Paragraph("REFERENCIAS PERSONALES", styles["h1"]))
    for ref in REFERENCES:
        story.append(Paragraph(ref["name"], styles["ref_name"]))
        story.append(
            Paragraph(f"{ref['role']} — {ref['org']} · {ref['phone']}", styles["ref_meta"])
        )

    doc.build(story)


def main() -> None:
    docx_path = OUT_DIR / "CV_ANDRES_ALFREDO_CORNEJO_FIGUEROA.docx"
    pdf_path = OUT_DIR / "CV_ANDRES_ALFREDO_CORNEJO_FIGUEROA.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"Wrote {docx_path}")
    print(f"Wrote {pdf_path}")


if __name__ == "__main__":
    main()
